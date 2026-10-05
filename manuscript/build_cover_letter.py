# -*- coding: utf-8 -*-
"""Cover letter for the submission to Superconductor Science and Technology (original submission); numbers from results.json.

    python build_cover_letter.py     ->  out/Cover_Letter_SST_v2.docx
"""
import json
import math
import os
import re
import sys

from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docx_helpers as H  # noqa: E402
from docx_helpers import new_document  # noqa: E402

R = json.load(open(os.path.join(HERE, "..", "results", "results.json"), encoding="utf-8"))
ZEN = json.load(open(os.path.join("C:" + os.sep, "YouTube", "_dxy_zenodo_state.json")))
SW, PP = ZEN.get("software_1.1.0", ZEN["software"])["doi"], ZEN.get("publication_v2", ZEN["publication"])["doi"]
REPO = "https://github.com/sandlerleon/damage-repair-xy-stiffness"
TH = R["thresholds"]
SL = R["timescale"]["slow_limit"]["5"]
F = R["fss"]["0.35"]["extrap_all"]
MP = R["mapping"]
BS = R["branches"]["summary"]
SWd = R["stiffness"]
TITLE = "Damage–repair kinetics and phase stiffness in a site-diluted XY model: implications for irradiated superconductors"


def br(mode):
    t = TH[mode]
    lo, hi = t["rho_vanished_max"], t["rho_finite_min"]
    return "%.2f–%.2f" % (lo, hi) if lo is not None else "below %.2f" % hi


under = int(5 * round(100 * max(1 - SWd["32|20|%s|%g|intact" % (m, r)]["ratio_eq_tw"] for m in ("keep", "neighbor") for r in (6.0, 7.0, 8.0, 10.0)
                                if SWd["32|20|%s|%g|intact" % (m, r)]["ratio_eq_tw"] is not None and SWd["32|20|%s|%g|intact" % (m, r)]["ups_tw"][0] > 0.2) / 5))
doc = new_document(size=11, line=1.15)


def para(text, bold=False, after=6):
    p = doc.add_paragraph()
    if bold:
        p.add_run(text).bold = True
    else:
        H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(3)


for line in ("Leon Sandler", "Independent researcher, Northbrook, Illinois, USA", "sandler.leon@gmail.com", "ORCID: https://orcid.org/0009-0007-4584-808X"):
    para(line, after=0)
para("")
para("4 October 2026")
para("The Editors\nSuperconductor Science and Technology", after=10)
para("Submission of a paper: “%s”" % TITLE, bold=True, after=10)
para("Dear Editors,")
para("I submit the enclosed paper for consideration in Superconductor Science and Technology. It asks what a minimal model can say about the "
     "competition between displacement damage and recovery in irradiated superconductors, the subject of recent work published in this journal on "
     "REBCO conductors for fusion magnets [Prokopec et al 2015, Fischer et al 2018, Unterrainer et al 2022, Adams et al 2023]. In those studies the "
     "transition temperature falls linearly with displacement damage, part of it is recovered by annealing, and the degradation is attributed to a loss of "
     "superfluid density. The paper takes the phase stiffness, the counterpart of the superfluid density, as its observable.")
para("Main results", bold=True, after=3)
bullet("A mean-field reduction of damage–repair kinetics, re-derived from the implemented update for the damaged fraction, predicts bistability. "
       "On the lattice, with up to %d seeds per condition and sizes to *L* = 64, no hysteresis and no dependence on the initial state is resolved (largest difference "
       "%.4f in the coherence, none significant after Holm correction; quasi-static ramps close), because the closure of the reduction fails: the multiplier "
       "of the damage rate that the lattice applies is a smooth function of the damaged fraction. The statement is made for the ranges tested." % (10, BS["max_abs_diff"]))
bullet("Without a field, and with repair that preserves the phase symmetry, repair restores spontaneous phase coherence and a stiffness, measured as the "
       "twist response of the non-equilibrium steady state with the winding number recorded, above a threshold ratio of repair to damage; finite-size scans (sizes to 96 and 64, operational cutoffs, not a "
       "thermodynamic-limit determination) place it at ρ = %s (repair that leaves the phase unchanged) and %s (repair to the phase of the neighbors), against about 2.5 without state-dependent damage. The "
       "equilibrium helicity-modulus formula underestimates this stiffness by up to about %d%% and is negative where the stiffness vanishes. The twist response is validated "
       "against the equilibrium helicity modulus on clean and diluted lattices, with imposed windings, and against an independent static code." % (br("keep"), br("neighbor"), under))
bullet("Numerically, the repair-to-damage ratio controls the state ever better as damage becomes slow compared with phase relaxation: three repair rules that differ by %.2f in the coherence "
       "at finite rates agree within %.3f at the slowest damage scanned (the limit itself is not claimed), and the equilibrium formula approaches the twist response." % (SL["spread_base"], SL["spread_slow"]))
bullet("Static random dilution of the lattice loses stiffness at a removed fraction of %.2f ± 0.01 (sizes to 96, finite-size extrapolation and its sensitivity "
       "reported); compact damage clusters tolerate considerably more. An illustrative effective mapping to displacement damage gives about %d model sites per dpa (a conversion factor, with no further physical interpretation), and an "
       "illustrative annealing analysis is compatible with a broad activation spectrum, whose width the two reported recoveries do not determine." %
       (round(F["f_inf"], 2), round(MP["sites_per_dpa"])))
para("The model is phenomenological and is not calibrated to any conductor; it has no pinning, so it says nothing about the critical current or its initial rise "
     "with fluence, and I have said so in the paper. The experimental values quoted are taken from the cited papers.", after=6)
para("Data, code and disclosures", bold=True, after=3)
para("All code, raw data and analysis scripts are public at %s, archived at https://doi.org/%s; the manuscript is available as a preprint at https://doi.org/%s. "
     "The simulation engine reproduces the update rules of the open-source MyUncle framework bit for bit. A separate manuscript on the same order-dependent-degradation "
     "mechanism in a rotor lattice without superconducting context (preprint https://doi.org/10.5281/zenodo.21210708) is intended for a different journal; the present paper does not depend on its conclusions. This manuscript "
     "has not been published and is not under consideration by any other journal. I declare no competing interests and no funding, and the use of AI-assisted tools is "
     "declared in the manuscript." % (REPO, SW, PP))
para("Thank you for considering the paper.")
para("Yours sincerely,", after=18)
para("Leon Sandler")
out = os.path.join(HERE, "out", "Cover_Letter_SST_v2.docx")
doc.save(out)
print("saved", out)
